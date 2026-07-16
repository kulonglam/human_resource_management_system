import csv
import io
from datetime import datetime

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _filename(prefix, ext):
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'{prefix}_{stamp}.{ext}'


def _safe_cell(value):
    """Prevent spreadsheet formula injection in user-controlled text."""
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return f"'{value}"
    return value


def _is_currency_column(column):
    key = column['key'].lower()
    return any(token in key for token in (
        'salary', 'allowance', 'deduction', 'tax', 'gross', 'net', 'amount', 'payroll',
    ))


def export_rows_csv(rows, columns, filename_prefix='report'):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c['label'] for c in columns])
    for row in rows or []:
        writer.writerow([_safe_cell(row.get(col['key'], '')) for col in columns])
    response = HttpResponse(f'\ufeff{buffer.getvalue()}', content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{_filename(filename_prefix, "csv")}"'
    return response


def export_rows_xlsx(rows, columns, filename_prefix='report', sheet_title='Report'):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    last_column = max(1, len(columns))
    generated_at = datetime.now().strftime('%d %b %Y, %H:%M')

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = ws.cell(row=1, column=1, value=sheet_title)
    title_cell.font = Font(size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill('solid', fgColor='004924')
    title_cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    ws.cell(row=2, column=1, value=f'Generated {generated_at} · Currency: UGX (Uganda Shilling)')
    ws.cell(row=2, column=1).font = Font(italic=True, color='404040')

    headers = [c['label'] for c in columns]
    keys = [c['key'] for c in columns]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='004924')
        cell.alignment = Alignment(vertical='center')

    for row in rows or []:
        ws.append([_safe_cell(row.get(key, '')) for key in keys])

    for index, column in enumerate(columns, start=1):
        if _is_currency_column(column):
            for cell in ws.iter_rows(min_row=5, min_col=index, max_col=index):
                cell[0].number_format = '"UGX" #,##0'
        values = [str(ws.cell(row=row, column=index).value or '') for row in range(4, ws.max_row + 1)]
        ws.column_dimensions[get_column_letter(index)].width = min(max(len(value) for value in values) + 3, 40)

    ws.freeze_panes = 'A5'
    if columns:
        ws.auto_filter.ref = f'A4:{get_column_letter(len(columns))}{ws.max_row}'
    ws.sheet_view.showGridLines = False

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{_filename(filename_prefix, "xlsx")}"'
    return response


def export_key_value_xlsx(data_dict, filename_prefix='report', sheet_title='Summary'):
    rows = [{'metric': key, 'value': value} for key, value in data_dict.items()]
    return export_rows_xlsx(
        rows,
        [
            {'key': 'metric', 'label': 'Metric'},
            {'key': 'value', 'label': 'Value'},
        ],
        filename_prefix,
        sheet_title,
    )


def export_rows_pdf(rows, columns, filename_prefix='report', title='Report'):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles['Title']),
        Paragraph(f'Generated {datetime.now().strftime("%d %b %Y %H:%M")} · Currency: UGX', styles['Normal']),
        Spacer(1, 12),
    ]
    table_data = [[col['label'] for col in columns]]
    for row in rows or []:
        table_data.append([_safe_cell(row.get(col['key'], '')) for col in columns])

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#004924')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    story.append(table)
    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{_filename(filename_prefix, "pdf")}"'
    return response
